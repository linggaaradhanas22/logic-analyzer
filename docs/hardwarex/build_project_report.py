"""Render the current Markdown report into an editable, styled A4 Word file."""
from pathlib import Path
import hashlib
import html
import json
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).with_name('manuscript.md')
OUT = ROOT / 'build/reports/project-documentation'
NAME = 'Logic Analyzer with Live Protocol Decoding_Lingga Aradhana Sahadewa'
INLINE = re.compile(r'(\*\*.+?\*\*|`.+?`|\*.+?\*)')


def runs(paragraph, text):
    for part in INLINE.split(text):
        run = paragraph.add_run(part.strip('`*'))
        if part.startswith('**'):
            run.bold = True
        elif part.startswith('*'):
            run.italic = True
        elif part.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        # Break long repository paths/URLs without changing their printed content.
        if len(part) > 70 and ('/' in part or '_' in part):
            run.text = run.text.replace('/', '/\u200b').replace('_', '_\u200b')


def parse(text):
    lines = text.splitlines()
    blocks = []
    i = 0
    start = re.compile(r'^(#{1,4} |\||!\[|- |\d+\. )')
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith('#'):
            hashes, title = line.split(' ', 1)
            blocks.append(('heading', (len(hashes), title)))
            i += 1
        elif line.startswith('!['):
            match = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', line)
            assert match, line
            blocks.append(('image', match.groups()))
            i += 1
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in row):
                    rows.append(row)
                i += 1
            blocks.append(('table', rows))
        else:
            kind = 'paragraph'
            if line.startswith('- '):
                kind, line = 'bullet', line[2:]
            elif re.match(r'^\d+\. ', line):
                kind = 'number'
            paragraphs = [line]
            i += 1
            while i < len(lines) and lines[i].strip() and not start.match(lines[i].strip()):
                paragraphs.append(lines[i].strip())
                i += 1
            blocks.append((kind, ' '.join(paragraphs)))
    return blocks


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), color)
    props.append(element)


def table_widths(head):
    widths = {2: [49, 123], 3: [40, 75, 57], 4: [33, 16, 21, 102],
              5: [12, 95, 15, 25, 25]}[len(head)]
    if head[0] == 'Week':
        return [14, 34, 109, 15]
    if head[0] == 'Class':
        return [30, 84, 58]
    if head[0] == 'Task':
        return [42, 18, 19, 93]
    if head[0] == 'Protocol':
        return [18, 26, 61, 67]
    return widths


def inline_html(text):
    output = []
    for part in INLINE.split(text):
        safe = html.escape(part.strip('`*'))
        if part.startswith('**'):
            safe = f'<strong>{safe}</strong>'
        elif part.startswith('*'):
            safe = f'<em>{safe}</em>'
        elif part.startswith('`'):
            safe = f'<code>{safe}</code>'
        output.append(safe)
    return ''.join(output)


def render_html(blocks):
    content = []
    for kind, value in blocks:
        if kind == 'heading':
            level, title = value
            content.append(f'<h{level}>{inline_html(title)}</h{level}>')
        elif kind == 'image':
            caption, path = value
            url = (SOURCE.parent / path).resolve().as_uri()
            content.append(f'<div class="figure"><img src="{url}" alt="{html.escape(caption)}"></div>')
        elif kind == 'table':
            head, *rows = value
            columns = '<colgroup>' + ''.join(f'<col style="width:{w / 172 * 100:.3f}%">' for w in table_widths(head)) + '</colgroup>'
            content.append('<table>' + columns + '<thead><tr>' + ''.join(f'<th>{inline_html(c)}</th>' for c in head) + '</tr></thead><tbody>')
            for row in rows:
                content.append('<tr>' + ''.join(f'<td>{inline_html(c)}</td>' for c in row) + '</tr>')
            content.append('</tbody></table>')
        elif kind in ['bullet', 'number']:
            content.append('<p class="item">' + ('• ' if kind == 'bullet' else '') + inline_html(value) + '</p>')
        else:
            cls = ' class="caption"' if value.startswith('*Figure ') else ''
            content.append(f'<p{cls}>{inline_html(value)}</p>')
    return '''<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Logic Analyzer with Live Protocol Decoding</title><style>
    body{font-family:Arial,Helvetica,sans-serif;font-size:10.5pt;line-height:1.25;color:#192B38;margin:0}
    h1,h2,h3{color:#173A56;break-after:avoid}h1{font-size:26pt;line-height:1.1;margin:0 0 7mm}
    h2{font-size:16pt;margin:7mm 0 3mm;border-bottom:1px solid #C8D7E3;padding-bottom:2mm}
    h3{font-size:12.5pt;margin:5mm 0 2.5mm}p{margin:0 0 3mm;orphans:3;widows:3;overflow-wrap:anywhere}
    code{font-family:Consolas,monospace;font-size:9pt}table{border-collapse:collapse;width:100%;table-layout:fixed;font-size:9pt;margin:3mm 0 5mm}
    th,td{border:1px solid #C9D3DB;text-align:left;vertical-align:top;padding:1.8mm;overflow-wrap:anywhere;line-height:1.2}
    th{background:#E6EDF3;font-weight:bold}thead{display:table-header-group}tr{break-inside:avoid}
    .figure{break-inside:avoid;break-after:avoid;margin:4mm 0 2mm}.figure img{width:100%;height:auto;display:block}
    .caption{font-size:9pt;break-before:avoid;margin-bottom:5mm;color:#425B6E}.item{margin-left:4mm}
    </style></head><body>''' + '\n'.join(content) + '</body></html>'


def main():
    assert OUT.is_dir()
    originals = OUT / 'originals'
    originals.mkdir(exist_ok=True)
    for filename in [NAME + '.docx', NAME + '.pdf']:
        origin = ROOT.parent / filename
        backup = originals / filename
        if not backup.exists():
            backup.write_bytes(origin.read_bytes())
    markdown = SOURCE.read_text(encoding='utf-8')
    blocks = parse(markdown)
    (OUT / 'project-report.html').write_text(render_html(blocks), encoding='utf-8')
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(18), Mm(18)
    section.left_margin, section.right_margin = Mm(19), Mm(19)
    section.header_distance, section.footer_distance = Mm(8), Mm(8)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        document.styles[name].font.name = 'Arial'
        document.styles[name].font.color.rgb = RGBColor.from_string('173A56')
    document.styles['Title'].font.size = Pt(26)
    document.styles['Heading 1'].font.size = Pt(17)
    document.styles['Heading 2'].font.size = Pt(13)
    document.styles['Heading 3'].font.size = Pt(11.5)
    header = section.header.paragraphs[0]
    header.add_run('LOGIC ANALYZER  |  FPGA + EMBEDDED SYSTEM').font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Lingga Aradhana Sahadewa  •  12 October 2026  |  ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    figures = []
    for kind, value in blocks:
        if kind == 'heading':
            level, title = value
            document.add_heading(title, level=max(0, level - 1))
        elif kind == 'image':
            caption, path = value
            image = (SOURCE.parent / path).resolve()
            assert image.is_file(), image
            p = document.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(image), width=Mm(172))
            figures.append(str(image))
        elif kind == 'table':
            head, *rows = value
            assert all(len(row) == len(head) for row in rows)
            table = document.add_table(rows=1, cols=len(head))
            table.style = 'Table Grid'
            header_row = table.rows[0]
            repeat = OxmlElement('w:tblHeader')
            header_row._tr.get_or_add_trPr().append(repeat)
            for cell, text in zip(header_row.cells, head):
                shade(cell, 'E6EDF3')
                runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    runs(cell.paragraphs[0], text)
            # Wide description columns, compact numeric columns.
            widths = table_widths(head)
            table.autofit = False
            for col, width in zip(table.columns, widths):
                col.width = Mm(width)
            for row in table.rows:
                no_split = OxmlElement('w:cantSplit')
                row._tr.get_or_add_trPr().append(no_split)
                for cell, width in zip(row.cells, widths):
                    cell.width = Mm(width)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.05
                        for run in p.runs:
                            run.font.size = Pt(9)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        else:
            style = {'bullet': 'List Bullet'}.get(kind)
            p = document.add_paragraph(style=style)
            runs(p, value)
            if value.startswith('*Figure '):
                p.paragraph_format.space_after = Pt(9)
                for run in p.runs:
                    run.font.size = Pt(9)
    document.core_properties.title = 'Logic Analyzer with Live Protocol Decoding'
    document.core_properties.author = 'Lingga Aradhana Sahadewa'
    document.core_properties.subject = 'Individual KCCI project for FPGA and Embedded system classes'
    destination = OUT / (NAME + '.docx')
    document.save(destination)
    reopened = Document(destination)
    assert len(reopened.inline_shapes) == 4
    abstract = markdown.split('## Abstract\n\n', 1)[1].split('\n\n## Keywords', 1)[0]
    assert len(abstract.split()) <= 200, len(abstract.split())
    assert '[Author' not in markdown and '[license]' not in markdown
    report = {'status': 'PASS', 'abstractWords': len(abstract.split()),
              'wordParagraphs': len(reopened.paragraphs), 'tables': len(reopened.tables),
              'images': len(reopened.inline_shapes), 'sourceSha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
              'figures': figures, 'output': str(destination)}
    (OUT / 'report-build.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
