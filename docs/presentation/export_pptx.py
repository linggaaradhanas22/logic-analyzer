"""Export browser-measured text/rectangles plus diagram images to editable PPTX."""
from pathlib import Path
import json
import os
import re

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "build" / "presentation"


def rgb(css):
    n = re.findall(r'\d+',css)
    return RGBColor(*map(int,n[:3]))


def main():
    layout = json.loads((OUT/'layout.json').read_text(encoding='utf-8'))
    content = json.loads((OUT/'slides.json').read_text(encoding='utf-8'))
    routes = json.loads((OUT/'routes.json').read_text(encoding='utf-8'))
    deck = Presentation()
    deck.slide_width = 12192000
    deck.slide_height = 6858000
    deck.core_properties.title = 'RTOS Logic Analyzer: Final Project Design'
    deck.core_properties.subject = 'ESP32-free direct-Zynq-USB architecture and verified results'
    deck.core_properties.author = os.environ.get('DECK_AUTHORS', 'Lingga Aradhana Sahadewa')
    deck.core_properties.keywords = 'Zynq,FPGA,FreeRTOS,USB,logic analyzer'
    jumps=[]
    for i,(g,s) in enumerate(zip(layout,content),1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        sx = deck.slide_width / g['width']
        sy = deck.slide_height / g['height']
        def geom(obj):
            return [int(obj[k]*factor) for k,factor in [('x',sx),('y',sy),('w',sx),('h',sy)]]
        for obj in g['rectangles']:
            sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,*geom(obj))
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(obj['fill']); sh.line.fill.background()
            # Remove theme-inherited shadows; the HTML has no shadow effects.
            style=sh._element.find(qn('p:style'))
            if style is not None:
                sh._element.remove(style)
        for obj in g['images']:
            slide.shapes.add_picture(str(OUT/obj['path']),*geom(obj))
        for obj in g['texts']:
            x,y,w,h = geom(obj)
            if obj.get('title'):
                w=int(deck.slide_width-x-58*sx)
            # Browser-measured lines remain independent editable text objects.
            shape = slide.shapes.add_textbox(x,y,max(w+int(sx*5),1),h+int(sy*3))
            shape.name = ('Title' if obj.get('title') else 'Text')+f' {i}: '+obj.get('text','')[:45]
            tf = shape.text_frame
            tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            tf.word_wrap=False
            p=tf.paragraphs[0]
            p.space_before=p.space_after=Pt(0)
            if obj.get('bullet'):
                p.text=' '
                pr=p._p.get_or_add_pPr()
                bullet=OxmlElement('a:buChar');bullet.set('char','•');pr.append(bullet)
            else:
                p.text=obj['text']
                pr=p._p.get_or_add_pPr();pr.append(OxmlElement('a:buNone'))
            for run in p.runs:
                run.font.name='Arial'
                run.font.size=Pt(obj['size']*0.75)
                run.font.color.rgb=rgb(obj['color'])
                run.font.bold=int(obj['weight'])>=600
                run.font.italic=obj.get('italic',False)
                if obj.get('href'):
                    run.hyperlink.address = obj['href']
                    run.font.underline = True
            if obj.get('text','').strip() in s.get('jumps',{}):
                jumps.append((shape,s['jumps'][obj['text'].strip()]))
            if obj.get('href'):
                shape.click_action.hyperlink.address = obj['href']
        slide.notes_slide.notes_text_frame.text = s['title']+'\n\n'+s['notes']+'\n\nSources: '+s['source']
    for shape,target in jumps:
        shape.click_action.target_slide=deck.slides[target-1]
    # Native PowerPoint custom shows skip the other class's detail section.
    shows=OxmlElement('p:custShowLst')
    for i,(name,indices) in enumerate(routes.items(),1):
        show=OxmlElement('p:custShow');show.set('name',name);show.set('id',str(i))
        entries=OxmlElement('p:sldLst')
        for index in indices:
            entry=OxmlElement('p:sld')
            entry.set(qn('r:id'),deck.slides._sldIdLst[index-1].get(qn('r:id')))
            entries.append(entry)
        show.append(entries);shows.append(show)
    default_style=deck._element.find(qn('p:defaultTextStyle'))
    if default_style is not None:
        default_style.addprevious(shows)
    else:
        deck._element.append(shows)
    dest=OUT/'logic-analyzer-design.pptx'
    deck.save(dest)
    # Full-slide review sheets; slide images are rendered HTML, not a PPT render.
    for start in range(0,len(layout),6):
        sheet=Image.new('RGB',(1536,3*452),'#E8EDF1')
        draw=ImageDraw.Draw(sheet)
        for k in range(start,min(start+6,len(layout))):
            im=Image.open(OUT/f'previews/slide-{k+1:02d}.png').convert('RGB')
            im.thumbnail((750,422))
            x=(k-start)%2*768+9;y=(k-start)//2*452+22
            sheet.paste(im,(x,y));draw.text((x,y-17),f'SLIDE {k+1}',fill='#173A56')
        sheet.save(OUT/f'previews/contact-{start//6+1}.png')
    # Inspect the saved package, including notes, dimensions and text editability.
    saved=Presentation(dest)
    assert len(saved.slides)==len(content)
    assert saved.slide_width==12192000 and saved.slide_height==6858000
    assert all(x.notes_slide.notes_text_frame.text for x in saved.slides)
    assert all(any(sh.has_text_frame and sh.text for sh in x.shapes) for x in saved.slides)
    assert len(saved._element.findall('.//'+qn('p:custShow')))==2
    print(f'PPTX package PASS: {len(saved.slides)} slides, 16:9, editable text, notes, 2 custom subsystem shows; {dest.stat().st_size:,} bytes')


if __name__=='__main__':
    main()
