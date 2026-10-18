"""Build the presentation speaking script into HTML, DOCX and a check report.

Source of truth: docs/presentation/presentation-script.md
Staging output:  build/presentation-script/
"""
from pathlib import Path
import json
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / 'docs' / 'presentation' / 'presentation-script.md'
OUT = ROOT / 'build' / 'presentation-script'
NAVY = '#173A56'
MUTED = '#607685'

REQUIRED_NUMBERS = ['1.62', '606,411', '606,405', '124.7', '32,768', '4,092', '67,108,864',
                    '41.4', '41.5', '37.93', '61.67', '0.576', '0.022', '18/18', '12/12',
                    '500', '180', '22/22', '8', '100', '256', '64', '0x40000000', '0x40400000',
                    '4C410002', '115,200', '10']
FORBIDDEN = ('consulting-pptx-skill', 'FINAL DIRECT-USB DESIGN', 'Sources:', 'Text 1', 'YYYY')

INLINE = re.compile(r'(\*\*.+?\*\*|\*.+?\*)')


def inline_html(text):
    out = []
    for part in INLINE.split(text):
        if part.startswith('**') and part.endswith('**'):
            out.append('<b>' + part[2:-2] + '</b>')
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            out.append('<i>' + part[1:-1] + '</i>')
        else:
            out.append(part)
    return ''.join(out)


def inline_runs(paragraph, text):
    for part in INLINE.split(text):
        if not part:
            continue
        run = paragraph.add_run(part.strip('*'))
        if part.startswith('**'):
            run.bold = True
        elif part.startswith('*'):
            run.italic = True


def parse(md_text):
    blocks = []
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('## '):
            blocks.append(('h2', line[3:].strip()))
            i += 1
        elif line.startswith('### '):
            blocks.append(('h3', line[4:].strip()))
            i += 1
        elif line.startswith('# '):
            blocks.append(('h1', line[2:].strip()))
            i += 1
        elif line.startswith('---'):
            blocks.append(('hr', ''))
            i += 1
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows))
        elif line.startswith('- '):
            items = []
            while i < len(lines) and lines[i].startswith('- '):
                items.append(lines[i][2:].strip())
                i += 1
            blocks.append(('bullets', items))
        elif line.startswith('> '):
            quote = []
            para = []
            while i < len(lines) and (lines[i].startswith('> ') or lines[i].strip() == '>'):
                if lines[i].strip() == '>':
                    if para:
                        quote.append(' '.join(para))
                        para = []
                else:
                    para.append(lines[i][2:].strip())
                i += 1
            if para:
                quote.append(' '.join(para))
            blocks.append(('quote', quote))
        else:
            para = [line.strip()]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r'^(#|\||- |> |---)', lines[i]):
                para.append(lines[i].strip())
                i += 1
            blocks.append(('para', ' '.join(para)))
    return blocks


def validate(blocks, md_text):
    slides = [b[1] for b in blocks if b[0] == 'h2' and b[1].startswith('Slide ')]
    numbers = []
    for title in slides:
        m = re.match(r'Slide (\d+) — ', title)
        assert m, f'Slide heading without number: {title}'
        numbers.append(int(m.group(1)))
    assert numbers == list(range(1, 39)), f'Slide numbers wrong: {numbers}'
    assert 'FPGA system' in md_text and 'Embedded system' in md_text
    missing = [n for n in REQUIRED_NUMBERS if n not in md_text]
    assert not missing, f'Missing key numbers: {missing}'
    bad = [f for f in FORBIDDEN if f in md_text]
    assert not bad, f'Forbidden text present: {bad}'
    # Readability: count sentences in spoken paragraphs only.
    spoken = ' '.join(' '.join(b[1]) for b in blocks if b[0] == 'quote')
    sentences = [s for s in re.split(r'(?<=[.!?])\s+', spoken) if len(s.split()) > 2]
    words = [len(s.split()) for s in sentences]
    stats = {'sentences': len(sentences), 'averageWords': round(sum(words) / len(words), 2),
             'longestSentence': max(words), 'over30Words': sum(1 for w in words if w > 30)}
    return {'slides': len(slides), 'spokenSentences': stats, 'keyNumbersPresent': len(REQUIRED_NUMBERS)}


def render_html(blocks, title):
    html = ['<!doctype html><html lang="en"><head><meta charset="utf-8">',
            f'<title>{title}</title><style>',
            'body{font-family:Arial,Helvetica,sans-serif;font-size:11pt;line-height:1.55;color:#233B4C;margin:0}',
            'main{max-width:180mm;margin:0 auto;padding:16mm 12mm}',
            'h1{font-size:22pt;color:%s;margin:0 0 6mm}' % NAVY,
            'h2{font-size:15pt;color:%s;border-bottom:2px solid %s;padding-bottom:2mm;margin:0 0 5mm;page-break-before:always}' % (NAVY, NAVY),
            'h2.first{page-break-before:avoid}',
            'h3{font-size:12.5pt;color:%s;margin:6mm 0 2mm}' % NAVY,
            'p{margin:0 0 3mm}',
            'blockquote{border-left:3px solid #C5A681;margin:0 0 4mm;padding:2mm 0 2mm 5mm;background:#F7F5F1}',
            'blockquote p{margin:0 0 2.5mm}',
            'ul{margin:0 0 3mm;padding-left:6mm}',
            'li{margin:0 0 1.6mm}',
            'table{border-collapse:collapse;width:100%;font-size:10pt;margin:0 0 4mm}',
            'th{text-align:left;border-bottom:2px solid %s;padding:2mm 2mm 2mm 0;color:%s}' % (NAVY, NAVY),
            'td{border-bottom:1px solid #D8DEE3;padding:1.8mm 2mm 1.8mm 0;vertical-align:top}',
            'hr{border:0;border-top:1px solid #C5A681;margin:6mm 0}',
            'em{color:%s}' % MUTED,
            '.foot{font-size:9pt;color:%s;margin-top:8mm}' % MUTED,
            '</style></head><body><main>']
    first_h2 = True
    for kind, value in blocks:
        if kind == 'h1':
            html.append(f'<h1>{inline_html(value)}</h1>')
        elif kind == 'h2':
            cls = ' class="first"' if first_h2 else ''
            first_h2 = False
            html.append(f'<h2{cls}>{inline_html(value)}</h2>')
        elif kind == 'h3':
            html.append(f'<h3>{inline_html(value)}</h3>')
        elif kind == 'para':
            html.append(f'<p>{inline_html(value)}</p>')
        elif kind == 'bullets':
            html.append('<ul>' + ''.join(f'<li>{inline_html(x)}</li>' for x in value) + '</ul>')
        elif kind == 'quote':
            html.append('<blockquote>' + ''.join(f'<p>{inline_html(p)}</p>' for p in value) + '</blockquote>')
        elif kind == 'table':
            head, *rows = value
            html.append('<table><thead><tr>' + ''.join(f'<th>{inline_html(c)}</th>' for c in head) + '</tr></thead><tbody>')
            for row in rows:
                html.append('<tr>' + ''.join(f'<td>{inline_html(c)}</td>' for c in row) + '</tr>')
            html.append('</tbody></table>')
        elif kind == 'hr':
            html.append('<hr>')
    html.append('</main></body></html>')
    return '\n'.join(html)


def render_docx(blocks, title, dest):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    for kind, value in blocks:
        if kind == 'h1':
            doc.add_heading(title, level=0)
        elif kind == 'h2':
            doc.add_page_break()
            doc.add_heading(value, level=1)
        elif kind == 'h3':
            doc.add_heading(value, level=2)
        elif kind == 'para':
            inline_runs(doc.add_paragraph(), value)
        elif kind == 'bullets':
            for item in value:
                inline_runs(doc.add_paragraph(style='List Bullet'), item)
        elif kind == 'quote':
            for para in value:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                p.paragraph_format.space_after = Pt(6)
                inline_runs(p, para)
        elif kind == 'table':
            head, *rows = value
            table = doc.add_table(rows=1, cols=len(head))
            table.style = 'Table Grid'
            for cell, text in zip(table.rows[0].cells, head):
                cell.paragraphs[0].add_run(text).bold = True
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.paragraphs[0].add_run(text)
        elif kind == 'hr':
            doc.add_paragraph('—' * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.author = 'Lingga Aradhana Sahadewa'
    doc.core_properties.title = title
    doc.save(dest)
    reopen = Document(dest)
    assert len(reopen.paragraphs) > 200
    return len(reopen.paragraphs)


def main():
    assert SRC.is_file(), 'Presentation script source is missing'
    assert (ROOT / 'build').is_dir(), 'Project build directory must exist'
    OUT.mkdir(exist_ok=True)
    md_text = SRC.read_text(encoding='utf-8')
    blocks = parse(md_text)
    report = validate(blocks, md_text)
    title = 'Presentation Script — RTOS Logic Analyzer'
    (OUT / 'presentation-script.html').write_text(render_html(blocks, title), encoding='utf-8')
    paragraphs = render_docx(blocks, title, OUT / 'presentation-script.docx')
    report['docxParagraphs'] = paragraphs
    report['htmlBytes'] = (OUT / 'presentation-script.html').stat().st_size
    (OUT / 'script-report.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
