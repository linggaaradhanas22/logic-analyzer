"""Check the actual presentation artifacts and retain verification reports."""
from pathlib import Path
import argparse
import hashlib
import json
import os
import subprocess
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw
import pymupdf
from publish_deck import main as publish

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'build/presentation'
SKILL=Path.home()/'.config/opencode/skills/consulting-pptx-skill'


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--check-only',action='store_true',help='Validate without publishing or cleaning staging')
    args=parser.parse_args()
    commands=[
        [sys.executable,'-X','utf8',str(SKILL/'scripts/check_deck.py'),str(OUT/'logic-analyzer-design.html')],
        [sys.executable,'-X','utf8',str(SKILL/'scripts/check_deck.py'),str(OUT/'logic-analyzer-design.pptx')],
        ['node',str(SKILL/'scripts/check_layout.mjs'),str(OUT/'logic-analyzer-design.html')],
    ]
    results=[]
    for label,cmd in zip(['check-html','check-pptx','check-layout'],commands):
        run=subprocess.run(cmd,cwd=ROOT,capture_output=True,text=True,encoding='utf-8',env=os.environ.copy())
        report=run.stdout+run.stderr
        (OUT/f'{label}.txt').write_text(report,encoding='utf-8')
        results.append({'check':label,'exitCode':run.returncode,'report':f'build/presentation/{label}.txt'})
        print(label+': '+report.strip().splitlines()[-1])
        if run.returncode: raise SystemExit(run.returncode)
    deck=Presentation(OUT/'logic-analyzer-design.pptx')
    slides=json.loads((OUT/'slides.json').read_text(encoding='utf-8'))
    routes=json.loads((OUT/'routes.json').read_text(encoding='utf-8'))
    count=len(slides)
    assert len(deck.slides)==count==38
    assert all(slide.notes_slide.notes_text_frame.text for slide in deck.slides)
    forbidden=('consulting-pptx-skill','final direct-usb design','sources:','fpga class','embedded systems class','embedded system class')
    for index,slide in enumerate(deck.slides,1):
        visible='\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame).lower()
        found=[label for label in forbidden if label in visible]
        assert not found,f'Removed label remains on slide {index}: {found}; text={visible}'
    assert all((OUT/f'previews/powerpoint-{i:02d}.png').is_file() for i in range(1,count+1))
    expected_names={'FPGA system':2,'Embedded system':3}
    assert set(routes)==set(expected_names)
    show_list=deck._element.find(qn('p:custShowLst'))
    assert show_list is not None
    actual_shows={show.get('name'):[entry.get(qn('r:id')) for entry in show.find(qn('p:sldLst'))] for show in show_list}
    assert set(actual_shows)==set(routes)
    for name,section in expected_names.items():
        path=routes[name]
        assert path==sorted(set(path))
        expected=[i for i,s in enumerate(slides,1) if s['category'] in (None,1,section,4)]
        assert path==expected,f'Cross-class detail or missing shared slides in {name}'
        rels=[deck.slides._sldIdLst[i-1].get(qn('r:id')) for i in path]
        assert actual_shows[name]==rels,f'PowerPoint custom show differs from route: {name}'
    for index,manifest in enumerate(slides):
        for label,target in manifest.get('jumps',{}).items():
            shapes=[s for s in deck.slides[index].shapes if s.has_text_frame and s.text.strip()==label]
            assert any(s.click_action.target_slide is not None and s.click_action.target_slide.slide_id==deck.slides[target-1].slide_id for s in shapes),f'Missing slide jump: {label}'
    report=json.loads((OUT/'render-report.json').read_text(encoding='utf-8'))
    assert not report['pageErrors'] and not report['geometryIssues']
    real_ui=ROOT/'docs/hardwarex/assets/uart-demo/uart-demo-capture.png'
    staged_ui=OUT/'uart-demo-capture.png'
    assert real_ui.is_file() and staged_ui.is_file(),'Real UART screenshot missing'
    assert hashlib.sha256(real_ui.read_bytes()).hexdigest()==hashlib.sha256(staged_ui.read_bytes()).hexdigest(),'Staged UART screenshot differs from source'
    html_source=(OUT/'logic-analyzer-design.html').read_text(encoding='utf-8')
    assert 'uart-demo-capture.png' in html_source,'Deck does not reference the real UART screenshot'
    frame=ROOT/'docs/hardwarex/assets/uart-demo/uart-demo-frame.png'
    assert hashlib.sha256(frame.read_bytes()).digest()==hashlib.sha256((OUT/frame.name).read_bytes()).digest(),'Staged decoding close-up differs from source'
    assert frame.name in slides[34]['body'],'Slide 35 is missing the decoding close-up'
    assert all(35 in route for route in routes.values()),'Decoding close-up missing from a class show'
    hardware=ROOT/'docs/hardwarex/assets/logic-analyzer-hardware.jpeg'
    assert hashlib.sha256(hardware.read_bytes()).digest()==hashlib.sha256((OUT/hardware.name).read_bytes()).digest(),'Staged hardware photo differs from source'
    assert hardware.name in slides[3]['body'],'Hardware photo missing from board-connections slide'
    video='https://youtu.be/1xx84Ndi4vQ?si=ldsuYEWvYVqu2u6-'
    assert any(run.hyperlink.address == video for shape in deck.slides[34].shapes
               if shape.has_text_frame for paragraph in shape.text_frame.paragraphs
               for run in paragraph.runs),'Video hyperlink missing from PPTX slide 35'
    with pymupdf.open(OUT/'logic-analyzer-design-powerpoint.pdf') as pdf:
        assert len(pdf)==count,'Native presentation PDF slide count differs'
        assert any(link.get('uri')==video for link in pdf[34].get_links()),'Video hyperlink missing from native PowerPoint PDF'
    ui=json.loads((OUT/'ui-screenshot-evidence.json').read_text(encoding='utf-8'))
    assert not ui['errors'],'Offline UI illustration produced page errors'
    original=json.loads((ROOT/'build/reports/presentation-delivery/comparison.json').read_text(encoding='utf-8'))
    for file in original['files']:
        comparison=ROOT.parent/'deliverables/comparison'/file['file']
        assert hashlib.sha256(comparison.read_bytes()).hexdigest()==file['sha256'],'Comparison copy changed'
    # Visual inspection sheets use the actual Microsoft PowerPoint renders.
    for start in range(0,count,6):
        sheet=Image.new('RGB',(1536,1356),'#E8EDF1');draw=ImageDraw.Draw(sheet)
        for i in range(start,min(start+6,count)):
            image=Image.open(OUT/f'previews/powerpoint-{i+1:02d}.png').convert('RGB')
            image.thumbnail((750,422))
            x=(i-start)%2*768+9;y=(i-start)//2*452+22
            sheet.paste(image,(x,y));draw.text((x,y-17),f'SLIDE {i+1}',fill='#173A56')
        sheet.save(OUT/f'previews/powerpoint-contact-{start//6+1}.png')
    (OUT/'delivery-checks.json').write_text(json.dumps({'slides':count,'subsystemRoutes':routes,'customShowsVerified':True,'powerPointNativeRender':True,'editableText':True,'notesPerSlide':True,'requestedVisibleRemovalsVerified':list(forbidden),'checks':results,'hardwareRetested':False},indent=2),encoding='utf-8')
    print(f'Delivery package PASS: {count} slides, 2 independent custom shows, notes, native renders, zero geometry issues; comparison intact; requested visible labels absent.')
    if not args.check_only:
        publish()


if __name__=='__main__':
    main()
